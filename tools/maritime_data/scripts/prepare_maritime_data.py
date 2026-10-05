"""Read-only source ingestion; emit a separate, reproducible Maritime data review package.

No database writes, fabricated observations, inferred joins, or official CII results.
Requires openpyxl for XLSX input. CSV-only parsing uses the Python standard library.
"""
import argparse
import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

VERSION = 'maritime_data-2026-09-30-v1'
MISSING = {'': 'blank', 'n/a': 'not_applicable', 'not applicable': 'not_applicable',
           'division by zero!': 'source_division_by_zero', 'nan': 'source_nan'}
FIELDS = {
    'source_files': [('source_id','text','PK','Dataset ID + SHA256'),('dataset_id','text','required','DS-001..006'),('filename','text','required','Original basename'),('sha256','text','required','Original byte checksum'),('source_url','text','nullable','Source recorded in assessment; not license approval'),('data_origin','text','required','REAL / SYNTHETIC / REFERENCE / SCHEMA'),('granularity','text','required','report / daily / position / port / schema'),('selection','text','required','Selected use for this preparation'),('row_count','integer','required','Count read from original'),('header_json','jsonb','required','Original ordered headers; duplicate XLSX names preserved')],
    'vessels': [('vessel_id','text','PK','REAL:IMO:<number> or SYN:<source vessel ID>'),('data_origin','text','required','REAL or SYNTHETIC; never joined across origins'),('imo_number','text','nullable','Seven-digit source IMO; synthetic remains NULL'),('source_vessel_id','text','required','Original IMO or synthetic vessel ID')],
    'annual_reports': [('record_id','text','PK','Stable source/sheet/row hash'),('source_id','text','FK source_files','Source version'),('source_sheet','text','required','CSV or original XLSX sheet'),('source_row','integer','required','Original row number; CSV is logical record number'),('vessel_id','text','FK vessels','Real vessel only'),('vessel_name','text','nullable','Name at reporting time'),('ship_type','text','nullable','Source classification, unmodified'),('reporting_year','integer','required','Reporting year'),('report_type','text','required','FULL / PARTIAL / ARCHIVE_ANNUAL'),('period_label','text','required','Original reporting period'),('period_start','date','nullable','Reported partial bounds or nominal annual boundary; not observation date'),('period_end','date','nullable','Reported partial bounds or nominal annual boundary'),('fuel_t','numeric','nullable','Source total fuel, tonnes; no fuel split invented'),('co2_t','numeric','nullable','Reported CO2 tonnes; not CO2-equivalent'),('sea_hours','numeric','nullable','Reported sea hours, aliases checked for agreement'),('fuel_kg_per_nm','numeric','nullable','Reported distance intensity'),('distance_nm_estimate','numeric','nullable','fuel_t * 1000 / intensity; scope not verified'),('dwt_t','numeric','nullable','Source value only; missing stays NULL'),('fuel_type','text','nullable','Not provided in MRV inputs'),('quality_status','text','required','VALID / REVIEW / REJECT'),('aggregate_eligible','boolean','required','False for partial, duplicate keys, invalid core numbers'),('provenance_json','jsonb','required','Per-field original columns, missing reasons, distance method')],
    'synthetic_noon': [('record_id','text','PK','Stable source/row hash'),('source_id','text','FK source_files','Source version'),('source_row','integer','required','Original CSV logical row'),('vessel_id','text','FK vessels','Synthetic vessel only'),('voyage_id','text','FK synthetic_voyages','Synthetic voyage key'),('report_date','date','required','Noon reporting date; observation period may straddle years'),('period_start_utc','timestamptz','required','Original interval start UTC'),('observed_at_utc','timestamptz','required','Original interval end UTC'),('vessel_name','text','required','Fictional source name'),('ship_type','text','required','Fictional source type'),('dwt_t','numeric','required','Fictional DWT tonnes'),('operating_status','text','required','Source SEA/PORT classification'),('distance_nm','numeric','required','Fictional daily distance; zero in port allowed'),('fuel_t','numeric','required','Fictional daily fuel tonnes'),('fuel_type','text','required','Source fuel label; not an official factor mapping'),('speed_kn','numeric','required','Fictional speed'),('engine_hours','numeric','required','0..24 hours'),('quality_status','text','required','VALID / REVIEW / REJECT'),('provenance_json','jsonb','required','Simulation flags/version/seed and original missing reasons')],
    'synthetic_voyages': [('voyage_id','text','PK','SYN:<source voyage ID>'),('vessel_id','text','FK vessels','Synthetic vessel'),('departure_port_label','text','nullable','Source label; no unverified WPI match'),('arrival_port_label','text','nullable','Source label; no unverified WPI match'),('first_report_date','date','required','First included report date, NOT departure time'),('last_report_date','date','required','Last included report date, NOT arrival time'),('mapping_status','text','required','UNMAPPED_PORT_LABELS / CONFLICT')],
    'quality_issues': [('issue_id','integer','PK','Sequence within deterministic run'),('record_id','text','nullable','Logical reference to annual/noon row'),('source_id','text','FK source_files','Input source version'),('source_row','integer','required','Original record row'),('severity','text','required','INFO / REVIEW / REJECT'),('field','text','required','Affected field'),('code','text','required','Machine-readable issue code'),('detail','text','required','Evidence without replacing original value')],
}

def text(value):
    if isinstance(value,float) and value.is_integer(): return str(int(value))
    return '' if value is None else str(value).strip()

def number(value):
    raw = text(value)
    if raw.lower() in MISSING: return None, MISSING[raw.lower()]
    # Commas are deliberately not guessed as decimal or thousands separators.
    try:
        n = Decimal(raw)
        if not n.is_finite(): return None, 'non_finite'
        if n < 0: return None, 'negative'
        return n, None
    except InvalidOperation: return None, 'invalid_numeric'

def resolve_hours(row):
    keys = ['annual_time_at_sea_h_reported','raw__Annual Total time spent at sea [hours]',
            'raw__Annual Time spent at sea [hours]','raw__Total time spent at sea [hours]',
            'raw__Time spent at sea [hours]']
    found = [(key, number(row.get(key))[0]) for key in keys]
    found = [(key, value) for key, value in found if value is not None]
    if not found: return None, {'reason':'no_valid_source_value'}
    if len({value for _,value in found}) != 1:
        return None, {'reason':'conflicting_source_values','values':{k:str(v) for k,v in found}}
    return found[0][1], {'columns':[k for k,_ in found], 'recovered_from_raw':found[0][0]!=keys[0]}

def period(label):
    match = re.fullmatch(r'(\d{4})(?:\s*\((\d{1,2})/(\d{1,2})\s*-\s*(\d{1,2})/(\d{1,2})\))?',text(label))
    if not match: return None, None, None
    y = int(match[1])
    try:
        start = date(y,int(match[3]),int(match[2])) if match[2] else date(y,1,1)
        end = date(y,int(match[5]),int(match[4])) if match[4] else date(y,12,31)
        if end < start: return y,None,None
        return y,str(start),str(end)
    except ValueError: return y,None,None

def digest(path):
    with path.open('rb') as f: return hashlib.file_digest(f,'sha256').hexdigest()

def csv_rows(path):
    for encoding in ['utf-8-sig','cp949']:
        try:
            with path.open(encoding=encoding,newline='') as f:
                reader=csv.DictReader(f); rows=list(reader); return reader.fieldnames,rows,encoding
        except UnicodeDecodeError: continue
    raise ValueError('Unsupported encoding: '+path.name)

def emit(path, rows, fields):
    with path.open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=fields); writer.writeheader()
        for row in rows:
            writer.writerow({key:json.dumps(value,ensure_ascii=False,default=str) if isinstance(value,(dict,list)) else ('true' if value else 'false') if isinstance(value,bool) else value for key,value in row.items() if key in fields})

def prepare(source, out):
    import openpyxl
    def locate(name):
        matches=list(source.rglob(name))
        if len(matches)!=1:raise ValueError(f'Expected exactly one {name}; found {len(matches)}')
        return matches[0]
    out.mkdir(parents=True,exist_ok=False)  # Never overwrite a prior run.
    tables={name:[] for name in FIELDS}; vessels={}; voyages={}; raw_path=out/'raw_records.jsonl'
    issues=tables['quality_issues']; annual=tables['annual_reports']; noon=tables['synthetic_noon']
    def issue(r, severity, field, code, detail):
        issues.append(dict(issue_id=len(issues)+1,record_id=r.get('record_id'),source_id=r['source_id'],source_row=r.get('source_row',0),severity=severity,field=field,code=code,detail=detail))
        if 'quality_status' in r and severity!='INFO':
            if r['quality_status']!='REJECT': r['quality_status']=severity
    def register(dataset, path, origin, granularity, selection, url, headers, count):
        sha=digest(path); sid=dataset+':'+sha
        tables['source_files'].append(dict(source_id=sid,dataset_id=dataset,filename=path.name,sha256=sha,source_url=url,data_origin=origin,granularity=granularity,selection=selection,row_count=count,header_json=headers))
        return sid
    def identity(sid,sheet,row): return hashlib.sha256(f'{sid}|{sheet}|{row}'.encode()).hexdigest()
    def add_vessel(vid,origin,original):
        vessels.setdefault(vid,dict(vessel_id=vid,data_origin=origin,imo_number=original if origin=='REAL' else None,source_vessel_id=original))
    def annual_row(sid,sheet,n,original,values,kind,hours_meta):
        imo=text(values['imo']); y,start,end=period(values['period'])
        r=dict(record_id=identity(sid,sheet,n),source_id=sid,source_sheet=sheet,source_row=n,vessel_id='REAL:IMO:'+imo,vessel_name=text(values['name']),ship_type=text(values['type']),reporting_year=y,report_type=kind,period_label=text(values['period']),period_start=start,period_end=end,fuel_type=None,quality_status='VALID',aggregate_eligible=True)
        prov={'data_origin':'REAL','granularity':'reported_period','period_basis':'nominal_year_bounds' if re.fullmatch(r'\d{4}',text(values['period'])) else 'source_period_label','missing':{'fuel_type':'not_provided'},'columns':values['columns'],'sea_hours':hours_meta,'distance':{'method':'fuel_t * 1000 / fuel_kg_per_nm','basis':'DERIVED_ESTIMATE','scope_alignment':'UNVERIFIED','independent_validation_target':False}}
        for key in ['fuel_t','co2_t','fuel_kg_per_nm','dwt_t','sea_hours']:
            v,reason=number(values.get(key));r[key]=v
            if reason: prov['missing'][key]=reason
            if reason in {'negative','invalid_numeric','non_finite'}:issue(r,'REVIEW',key,reason,text(values.get(key)))
        r['distance_nm_estimate']=r['fuel_t']*1000/r['fuel_kg_per_nm'] if r['fuel_t'] is not None and r['fuel_kg_per_nm'] not in (None,0) else None
        if r['distance_nm_estimate'] is None:prov['missing']['distance_nm_estimate']='missing_numerator_or_missing_zero_denominator'
        if not re.fullmatch(r'\d{7}',imo):issue(r,'REJECT','imo_number','invalid_imo_format',imo)
        else:
            if sum(int(imo[i])*(7-i) for i in range(6))%10!=int(imo[-1]):issue(r,'REVIEW','imo_number','imo_checksum_mismatch',imo)
        if not start:issue(r,'REJECT','period','invalid_period',r['period_label'])
        if hours_meta.get('reason')=='conflicting_source_values':issue(r,'REVIEW','sea_hours','alias_conflict',json.dumps(hours_meta))
        if hours_meta.get('recovered_from_raw'):issue(r,'INFO','sea_hours','recovered_from_raw',','.join(hours_meta['columns']))
        if r['fuel_t'] is None or r['co2_t'] is None:issue(r,'REVIEW','totals','missing_core_total','No imputation')
        if kind=='PARTIAL':issue(r,'REVIEW','report_type','partial_report','Keep separately; do not add to Full reports')
        r['provenance_json']=prov; add_vessel(r['vessel_id'],'REAL',imo); annual.append(r)
        return r

    with raw_path.open('w',encoding='utf-8') as raw:
        path=locate('eu_mrv_vessel_annual_2018_2022.csv'); headers,rows,_=csv_rows(path)
        sid=register('DS-005',path,'REAL','report','REAL_ANNUAL_QUERY','https://github.com/richhuwtaylor/ship-emissions',headers,len(rows))
        for n,d in enumerate(rows,2):
            hours,hmeta=resolve_hours(d)
            cols={'imo':'imo_number','name':'vessel_name','type':'ship_type','period':'reporting_year','fuel_t':'fuel_consumption_t_reported','co2_t':'co2_emissions_t_reported','fuel_kg_per_nm':'fuel_per_distance_kg_per_nm_reported','dwt_t':'dwt_t'}
            v={k:d.get(c) for k,c in cols.items()};v.update(sea_hours=hours,columns=cols)
            r=annual_row(sid,'CSV',n,d,v,'ARCHIVE_ANNUAL',hmeta)
            original_distance,_=number(d.get('distance_nm_derived'))
            if original_distance is not None and r['distance_nm_estimate'] is not None and abs(original_distance-r['distance_nm_estimate'])>Decimal('.001'):
                issue(r,'REVIEW','distance_nm_estimate','distance_recompute_mismatch',str(original_distance))
            raw.write(json.dumps({'record_id':r['record_id'],'original':d},ensure_ascii=False)+'\n')

        path=locate('2025-v57-12092026-EU_MRV_Publication_of_information.xlsx')
        wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
        sheet_rows={};sheet_headers={}
        for s in wb:
            sheet_headers[s.title]=list(next(s.iter_rows(min_row=3,max_row=3,values_only=True)))
            sheet_rows[s.title]=[(n,list(row)) for n,row in enumerate(s.iter_rows(min_row=4,values_only=True),4) if any(v is not None for v in row)]
        sid=register('DS-001',path,'REAL','report','REAL_ANNUAL_QUERY','https://mrv.emsa.europa.eu',sheet_headers,sum(map(len,sheet_rows.values())))
        for sheet,rows in sheet_rows.items():
            headers=sheet_headers[sheet]
            if headers[:4]!=['IMO Number','Name','Ship type','Reporting Period']:raise ValueError('Unexpected ship header layout')
            def unique_col(label):
                if headers.count(label)!=1:raise ValueError('Missing/ambiguous numeric header: '+label)
                return headers.index(label)
            cols={'fuel_t':unique_col('Total fuel consumption [m tonnes]'),'co2_t':unique_col('Total CO₂ emissions [m tonnes]'),'sea_hours':unique_col('Time spent at sea [hours]'),'fuel_kg_per_nm':unique_col('Fuel consumption per distance [kg / n mile]')}
            if sheet not in {'2025 Full ERs','2025 Partial ERs'}:raise ValueError('Unexpected report sheet: '+sheet)
            for n,values in rows:
                v={'imo':values[0],'name':values[1],'type':values[2],'period':values[3],'columns':{'imo':'A (Ship)','name':'B (Ship)',**{k:i+1 for k,i in cols.items()}},**{k:values[i] for k,i in cols.items()}}
                r=annual_row(sid,sheet,n,values,v,'PARTIAL' if 'Partial' in sheet else 'FULL',{'columns':['Time spent at sea [hours]']})
                raw.write(json.dumps({'record_id':r['record_id'],'original_values':values},ensure_ascii=False,default=str)+'\n')
        wb.close()

        path=locate('synthetic_noon_daily_2025.csv');headers,rows,_=csv_rows(path)
        sid=register('DS-006',path,'SYNTHETIC','daily','DEVELOPMENT_ONLY','',headers,len(rows))
        for n,d in enumerate(rows,2):
            vid='SYN:'+d['vessel_id']; voyage='SYN:'+d['voyage_id']
            r=dict(record_id=identity(sid,'CSV',n),source_id=sid,source_row=n,vessel_id=vid,voyage_id=voyage,report_date=d['report_date'],period_start_utc=d['period_start_utc'],observed_at_utc=d['observed_at_utc'],vessel_name=d['vessel_name'],ship_type=d['ship_type'],operating_status=d['operating_status'],fuel_type=d['fuel_type'],quality_status='VALID')
            prov={'data_origin':d['data_origin'],'granularity':'daily_interval','simulation_version':d['simulation_version'],'seed':d['synthetic_seed'],'source_quality_flag':d['quality_flag'],'missing':{},'columns':{},'latitude':None,'longitude':None,'imo_number':None}
            for dest,src in [('dwt_t','dwt_t'),('distance_nm','distance_nm'),('fuel_t','fuel_consumed_t'),('speed_kn','speed_knots'),('engine_hours','engine_hours')]:
                r[dest],reason=number(d[src]);prov['columns'][dest]=src
                if reason:prov['missing'][dest]=reason;issue(r,'REJECT',dest,reason,d[src])
            if d['data_origin']!='SYNTHETIC' or d['imo_number']:issue(r,'REJECT','data_origin','unexpected_identity','Synthetic namespace requires SYNTHETIC and blank IMO')
            if r['dwt_t']==0:issue(r,'REJECT','dwt_t','zero_capacity','Must be positive')
            try:
                date.fromisoformat(r['report_date']);start=datetime.fromisoformat(r['period_start_utc'].replace('Z','+00:00'));end=datetime.fromisoformat(r['observed_at_utc'].replace('Z','+00:00'))
                if start.tzinfo is None or end.tzinfo is None or not 0<(end-start).total_seconds()<=86400:raise ValueError()
                if str(end.date())!=r['report_date']:issue(r,'REVIEW','report_date','date_interval_mismatch','Date differs from observation end')
            except ValueError:issue(r,'REJECT','period','invalid_interval','Expected valid date and bounded UTC interval')
            if r['engine_hours'] is not None and r['engine_hours']>24:issue(r,'REJECT','engine_hours','out_of_range','Above 24 hours')
            if d['operating_status']=='SEA' and all(r[k] is not None for k in ['distance_nm','speed_kn','engine_hours']):
                if abs(r['distance_nm']-r['speed_kn']*r['engine_hours'])>Decimal('.15'):issue(r,'REVIEW','distance_nm','simulation_distance_mismatch','Distance differs from speed * engine hours by > 0.15 nm')
            if d['operating_status']=='PORT' and r['distance_nm']!=0:issue(r,'REVIEW','distance_nm','port_distance','Source PORT row has nonzero distance')
            r['provenance_json']=prov;noon.append(r);add_vessel(vid,'SYNTHETIC',d['vessel_id'])
            v=voyages.setdefault(voyage,dict(voyage_id=voyage,vessel_id=vid,departure_port_label=d['departure_port'],arrival_port_label=d['arrival_port'],first_report_date=r['report_date'],last_report_date=r['report_date'],mapping_status='UNMAPPED_PORT_LABELS'))
            if (v['vessel_id'],v['departure_port_label'],v['arrival_port_label'])!=(vid,d['departure_port'],d['arrival_port']):
                v['mapping_status']='CONFLICT';issue(r,'REVIEW','voyage_id','voyage_metadata_conflict',voyage)
            v['first_report_date']=min(v['first_report_date'],r['report_date']);v['last_report_date']=max(v['last_report_date'],r['report_date'])
            raw.write(json.dumps({'record_id':r['record_id'],'original':d},ensure_ascii=False)+'\n')

    for rows,keys in [(annual,['vessel_id','reporting_year','report_type']),(noon,['vessel_id','report_date'])]:
        counts=Counter(tuple(r[k] for k in keys) for r in rows)
        for r in rows:
            if counts[tuple(r[k] for k in keys)]>1:issue(r,'REVIEW',','.join(keys),'duplicate_business_key','All source rows retained; no arbitrary winner')
    for r in annual:r['aggregate_eligible']=r['quality_status']=='VALID' and r['report_type']!='PARTIAL'
    for dataset,name,origin,granularity,selection,url in [
        ('DS-002','해양수산부_선박_AIS_동적정보_20220101.csv','REAL','position','REFERENCE_ONLY_MASKED_ID','https://www.data.go.kr/'),
        ('DS-003','UpdatedPub150.csv','REFERENCE','port','PORT_REFERENCE_PENDING_MAPPING','https://msi.nga.mil/Publications/WPI')]:
        path=locate(name);headers,rows,encoding=csv_rows(path);register(dataset,path,origin,granularity,selection,url,headers,len(rows))
    path=locate('Smart-Maritime-Council-Standardised-Vessel-Dataset-SVD-for-Noon-Reports-and-Emissions-Reporting-V2-May-2025.xlsx')
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    inventory={s.title:sum(1 for r in s.iter_rows(values_only=True) if any(v is not None for v in r)) for s in wb};wb.close()
    register('DS-004',path,'SCHEMA','schema','SCHEMA_REFERENCE_ONLY','https://smartmaritimenetwork.com/2025/05/28/smart-maritime-council-adds-emissions-data-to-standardised-vessel-dataset-with-version-2-0-update/',inventory,sum(inventory.values()))
    tables['vessels']=list(vessels.values());tables['synthetic_voyages']=list(voyages.values())
    for name,rows in tables.items():emit(out/(name+'.csv'),rows,[x[0] for x in FIELDS[name]])
    # Every normalized observation must remain traceable, including REVIEW/REJECT.
    expected=sum(s['row_count'] for s in tables['source_files'] if s['dataset_id'] in {'DS-001','DS-005','DS-006'})
    assert len(annual)+len(noon)==expected
    assert all(r['vessel_id'] in vessels for r in annual+noon)
    assert all(r['voyage_id'] in voyages for r in noon)
    assert len({r['record_id'] for r in annual+noon})==expected
    summary={'pipeline_version':VERSION,'rows':{k:len(v) for k,v in tables.items()},'annual_report_types':dict(Counter(r['report_type'] for r in annual)),
             'annual_quality':dict(Counter(r['quality_status'] for r in annual)),'noon_quality':dict(Counter(r['quality_status'] for r in noon)),
             'annual_aggregate_eligible':sum(r['aggregate_eligible'] for r in annual),'issues':dict(Counter(i['code'] for i in issues)),
             'missing_annual':{k:sum(r[k] is None for r in annual) for k in ['fuel_t','co2_t','sea_hours','dwt_t','distance_nm_estimate']},
             'synthetic_zero_distance_rows':sum(r['distance_nm']==0 for r in noon),'sources':tables['source_files'],
             'checks':['observation count reconciled','unique source record IDs','all vessel/voyage references resolved'],
             'limits':['No official CII result','No DB load performed','Port labels not joined to WPI','AIS not joined to real or synthetic vessels','License terms not verified','Derived distance not an independent validation target']}
    (out/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    (out/'dictionary.json').write_text(json.dumps(FIELDS,ensure_ascii=False,indent=2),encoding='utf-8')
    return summary

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--source',required=True,type=Path);p.add_argument('--out',required=True,type=Path);a=p.parse_args()
    result=prepare(a.source,a.out)
    print(json.dumps({k:v for k,v in result.items() if k!='sources'},ensure_ascii=False,indent=2))
