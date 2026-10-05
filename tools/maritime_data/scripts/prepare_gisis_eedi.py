#!/usr/bin/env python3
"""Normalize the anonymized IMO GISIS EEDI workbook into a reference CSV.

The source does not contain IMO numbers. Records produced here must never be
joined to maritime_data.vessels as if they were identified ships.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SOURCE_URL = "https://gisis.imo.org/Public/MARPOL6/EEDIData.aspx"
SOURCE_UPDATED = "2025-12-02"

SHEETS = {
    "Bulk": ("Bulk carrier", "DWT"),
    "Gas": ("Gas carrier", "DWT"),
    "Tanker": ("Tanker", "DWT"),
    "Container": ("Containership", "DWT"),
    "General": ("General cargo ship", "DWT"),
    "Ref.": ("Refrigerated cargo carrier", "DWT"),
    "Comb.": ("Combination carrier", "DWT"),
    "LNG": ("LNG carrier", "DWT"),
    "Ro-ro (VC)": ("Ro-ro cargo ship (vehicle carrier)", "DWT"),
    "Ro-ro cargo": ("Ro-ro cargo ship", "DWT"),
    "RoPass.": ("Ro-ro passenger ship", "DWT"),
    "Cruise": ("Cruise passenger ship, non-conventional propulsion", "GT"),
}

FIELDNAMES = [
    "reference_id",
    "source_sheet",
    "source_row",
    "anonymized_ship_no",
    "ship_type",
    "applicable_phase",
    "capacity",
    "capacity_unit",
    "year_of_delivery",
    "length_between_perpendiculars_m",
    "breadth_m",
    "draught_m",
    "required_eedi_mandatory",
    "required_eedi_non_mandatory",
    "required_eedi",
    "attained_eedi_mandatory",
    "attained_eedi_non_mandatory",
    "attained_eedi",
    "vref_kn",
    "main_engine_power_kw",
    "reduction_rate_reference",
    "reduction_rate_previous_reference",
    "fuel_type",
    "fdf_gas",
    "ice_class",
    "common_commercial_size",
    "innovative_electrical_technology",
    "innovative_mechanical_technology",
    "design_statement",
    "quality_status",
    "source_updated",
    "source_url",
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def number(value: Any) -> int | float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer():
            return int(value)
        return value
    return None


def positive(value: Any) -> int | float | None:
    parsed = number(value)
    return parsed if parsed is not None and parsed > 0 else None


def first_positive(*values: Any) -> int | float | None:
    for value in values:
        parsed = positive(value)
        if parsed is not None:
            return parsed
    return None


def slug(value: str) -> str:
    return re.sub(r"[^A-Z0-9]+", "-", value.upper()).strip("-")


def normalize_row(sheet_name: str, source_row: int, raw: tuple[Any, ...]) -> dict[str, Any]:
    ship_type, capacity_unit = SHEETS[sheet_name]
    expanded = list(raw) + [None] * max(0, 24 - len(raw))
    ro_variant = sheet_name in {"Ro-ro cargo", "RoPass."}
    mandatory_required = positive(expanded[7])
    voluntary_required = positive(expanded[8])
    mandatory_attained = positive(expanded[9])
    voluntary_attained = positive(expanded[10])
    capacity = positive(expanded[2])
    year = positive(expanded[6])
    attained = first_positive(mandatory_attained, voluntary_attained)

    if ro_variant:
        previous_reduction = number(expanded[17])
        reference_reduction = number(expanded[18])
        commercial_size, fuel_type, fdf_gas, ice_class, statement = expanded[19:24]
    else:
        previous_reduction = None
        reference_reduction = number(expanded[17])
        commercial_size, fuel_type, fdf_gas, ice_class, statement = expanded[18:23]

    issues = []
    if capacity is None:
        issues.append("missing_capacity")
    if attained is None:
        issues.append("missing_attained_eedi")
    if year is None:
        issues.append("missing_year_of_delivery")

    anonymized_no = int(expanded[1])
    return {
        "reference_id": f"GISIS-EEDI-2025:{slug(sheet_name)}:{anonymized_no}",
        "source_sheet": sheet_name,
        "source_row": source_row,
        "anonymized_ship_no": anonymized_no,
        "ship_type": ship_type,
        "applicable_phase": clean_text(expanded[0]),
        "capacity": capacity,
        "capacity_unit": capacity_unit,
        "year_of_delivery": year,
        "length_between_perpendiculars_m": positive(expanded[3]),
        "breadth_m": positive(expanded[4]),
        "draught_m": positive(expanded[5]),
        "required_eedi_mandatory": mandatory_required,
        "required_eedi_non_mandatory": voluntary_required,
        "required_eedi": first_positive(mandatory_required, voluntary_required),
        "attained_eedi_mandatory": mandatory_attained,
        "attained_eedi_non_mandatory": voluntary_attained,
        "attained_eedi": attained,
        "vref_kn": positive(expanded[11]),
        "main_engine_power_kw": positive(expanded[12]),
        "reduction_rate_reference": reference_reduction,
        "reduction_rate_previous_reference": previous_reduction,
        "fuel_type": clean_text(fuel_type),
        "fdf_gas": clean_text(fdf_gas),
        "ice_class": clean_text(ice_class),
        "common_commercial_size": clean_text(commercial_size),
        "innovative_electrical_technology": clean_text(expanded[14]) if clean_text(expanded[13]) == "Yes" else None,
        "innovative_mechanical_technology": clean_text(expanded[16]) if clean_text(expanded[15]) == "Yes" else None,
        "design_statement": clean_text(statement),
        "quality_status": "VALID" if not issues else "REVIEW:" + ",".join(issues),
        "source_updated": SOURCE_UPDATED,
        "source_url": SOURCE_URL,
    }


def prepare(source: Path, output_dir: Path) -> dict[str, Any]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    rows: list[dict[str, Any]] = []
    for sheet_name in SHEETS:
        worksheet = workbook[sheet_name]
        for source_row, raw in enumerate(worksheet.iter_rows(min_row=5, values_only=True), start=5):
            if not isinstance(raw[1], (int, float)):
                continue
            rows.append(normalize_row(sheet_name, source_row, raw))

    if len(rows) != 11244:
        raise ValueError(f"Expected 11244 anonymized ship rows, found {len(rows)}")
    if len({row['reference_id'] for row in rows}) != len(rows):
        raise ValueError("reference_id is not unique")

    output_dir.mkdir(parents=True, exist_ok=True)
    csv_path = output_dir / "eedi_reference_records.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDNAMES)
        writer.writeheader()
        writer.writerows(rows)

    sheet_counts = Counter(row["source_sheet"] for row in rows)
    status_counts = Counter(row["quality_status"] for row in rows)
    summary = {
        "prepared_on": date.today().isoformat(),
        "source": {
            "filename": source.name,
            "sha256": sha256(source),
            "url": SOURCE_URL,
            "updated": SOURCE_UPDATED,
            "access": "GISIS public account; anonymized and rounded EEDI database",
            "license": "No reuse licence stated on the download page; retain provenance and verify reuse terms before redistribution.",
        },
        "records": len(rows),
        "ship_type_counts": dict(sorted(Counter(row["ship_type"] for row in rows).items())),
        "sheet_counts": dict(sheet_counts),
        "capacity_units": dict(Counter(row["capacity_unit"] for row in rows)),
        "quality_status_counts": dict(status_counts),
        "completeness": {
            "capacity": sum(row["capacity"] is not None for row in rows),
            "year_of_delivery": sum(row["year_of_delivery"] is not None for row in rows),
            "attained_eedi": sum(row["attained_eedi"] is not None for row in rows),
            "vref_kn": sum(row["vref_kn"] is not None for row in rows),
            "main_engine_power_kw": sum(row["main_engine_power_kw"] is not None for row in rows),
            "fuel_type": sum(row["fuel_type"] is not None for row in rows),
        },
        "limitations": [
            "Records are anonymized and cannot be joined to maritime_data.vessels or annual_reports by IMO number.",
            "EEDI is a design-efficiency reference and is not annual operational DCS/CII input.",
            "Capacity, dimensions, speed, and power are rounded by the IMO Secretariat as described in the workbook explanatory note.",
            "The workbook does not supply annual fuel consumption or annual distance travelled.",
            "The dataset is retained as a separate reference set and is not loaded into maritime_data operational query views.",
        ],
    }
    summary_path = output_dir / "eedi_reference_summary.json"
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return summary


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output_dir", type=Path)
    args = parser.parse_args()
    print(json.dumps(prepare(args.source, args.output_dir), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
